import csv
import tempfile

import openpyxl

from core.report import export_events_excel


def test_export_produces_workbook():
    with tempfile.TemporaryDirectory() as d:
        csv_path = f"{d}/e.csv"
        with open(csv_path, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f)
            w.writerow(["时间", "类别", "最高置信度", "命中数量", "抓拍图片"])
            w.writerow(["2026-10-31 08:00:00", "head", 0.89, 3, "a.jpg"])
            w.writerow(["2026-10-31 09:00:00", "helmet", 0.88, 1, ""])
        out = export_events_excel(csv_path, f"{d}/o.xlsx")
        wb = openpyxl.load_workbook(out)
        assert set(wb.sheetnames) == {"事件明细", "统计汇总"}
        assert wb["事件明细"].max_row == 3            # 表头 + 2 条
